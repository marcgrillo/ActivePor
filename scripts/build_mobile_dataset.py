"""Build data/mobile from the spreadsheet of the experiment of Ciomek, Kadzinski & Tervonen (2017):
10 mobile-phone contracts evaluated on 4 criteria and the complete rankings stated by 94 decision makers.
The spreadsheet is not distributed with this repository; it is available from the authors of that study.

    python scripts/build_mobile_dataset.py path/to/mobile-phone-experiment.xlsx

Sheet 'Arkusz1' holds the contracts (fee and contract length are costs, data and roaming are gains);
sheet 'Arkusz2' holds one ranking per DM as 'i-j-k-...' (contract ids 0-9, best first).  The two cost
criteria are negated, so that all criteria are of gain type.  Each DM answers a pair from its own
ranking, and gets its own random question order (--seed makes the orders repeatable).
"""
import argparse
import os
import sys

import numpy as np

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
sys.path.insert(0, ROOT)
from activepor import data as D           # noqa: E402

COST = [0, 1]          # monthly fee (PLN), contract length (months)


def read(xlsx):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    rows = [r for r in wb['Arkusz1'].iter_rows(values_only=True) if isinstance(r[0], (int, float))]
    table = np.array([[r[1], r[2], r[3], r[4]] for r in rows], dtype=float)
    orders = [[int(x) for x in r[3].split('-')] for r in wb['Arkusz2'].iter_rows(values_only=True)
              if isinstance(r[0], (int, float)) and isinstance(r[3], str)]
    return table, np.array(orders)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('xlsx')
    ap.add_argument('--out', default=os.path.join(ROOT, 'data', 'mobile'))
    ap.add_argument('--seed', type=int, help='seed of the question orders (default: fresh randomness)')
    a = ap.parse_args()
    seed = int(np.random.SeedSequence(a.seed).entropy)       # recorded in meta.json
    rng = np.random.default_rng(seed)
    table, orders = read(a.xlsx)
    n, m = table.shape
    assert all(sorted(o) == list(range(n)) for o in orders)
    gain = table.copy()
    gain[:, COST] *= -1.0
    pairs = np.array([(i, j) for i in range(n) for j in range(i + 1, n)])
    ranking, answers = [], []
    for s, order in enumerate(orders):
        rank = np.empty(n, dtype=int)
        rank[order] = np.arange(n)
        perm = rng.permutation(len(pairs))
        answers.append([(i, j) if rank[i] < rank[j] else (j, i) for i, j in pairs[perm]])
        ranking.append(rank)
    H = len(orders)
    D.save_config(a.out, n, m, np.repeat(gain[None], H, axis=0), ranking, answers)
    D.save_meta(a.out, {'utility': 'real', 'answer_rule': 'stated ranking', 'instances': H, 'seed': seed,
                        'cost_criteria_negated': COST,
                        'source': 'Ciomek, Kadzinski & Tervonen (2017), Omega 71, Section 8'})
    print(f'wrote {a.out}: {H} DMs, {n} contracts, {m} criteria, {len(pairs)} pairs each')


if __name__ == '__main__':
    main()
